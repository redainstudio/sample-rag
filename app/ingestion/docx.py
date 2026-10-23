import io
from typing import List, Optional
import docx
from app.core.logging import logger
from app.ingestion.base import BaseParser
from app.schemas.document import ParsedSection


class DOCXParser(BaseParser):
    """Structure-aware DOCX parser supporting headings, paragraphs, and tables."""

    def parse_bytes(self, content: bytes, filename: str) -> List[ParsedSection]:
        sections: List[ParsedSection] = []
        try:
            doc = docx.Document(io.BytesIO(content))
            current_title = filename or "Document Overview"
            current_elements: List[str] = []

            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue

                style_name = paragraph.style.name.lower() if paragraph.style else ""

                # Check if paragraph is a heading
                if "heading" in style_name or style_name.startswith("title"):
                    if current_elements:
                        content_str = "\n".join(current_elements).strip()
                        if content_str:
                            sections.append(
                                ParsedSection(
                                    title=current_title,
                                    content=content_str,
                                    page_number=1,
                                    section_type="text",
                                    metadata={"filename": filename}
                                )
                            )
                        current_elements = []
                    current_title = text
                else:
                    current_elements.append(text)

            # Flush paragraphs
            if current_elements:
                content_str = "\n".join(current_elements).strip()
                if content_str:
                    sections.append(
                        ParsedSection(
                            title=current_title,
                            content=content_str,
                            page_number=1,
                            section_type="text",
                            metadata={"filename": filename}
                        )
                    )

            # Process tables: format each table into Markdown format to preserve headers
            for idx, table in enumerate(doc.tables):
                table_lines = []
                headers = []
                for row_idx, row in enumerate(table.rows):
                    row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    if row_idx == 0:
                        headers = row_cells
                        table_lines.append("| " + " | ".join(headers) + " |")
                        table_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
                    else:
                        table_lines.append("| " + " | ".join(row_cells) + " |")

                table_content = "\n".join(table_lines)
                if table_content:
                    sections.append(
                        ParsedSection(
                            title=f"Table {idx + 1} - {current_title}",
                            content=table_content,
                            page_number=1,
                            section_type="table",
                            metadata={"filename": filename, "table_index": idx + 1}
                        )
                    )

        except Exception as e:
            logger.error(f"Failed to parse DOCX {filename}: {e}")
            raise e

        return sections

    def parse_text(self, text: str, filename: Optional[str] = None) -> List[ParsedSection]:
        return [
            ParsedSection(
                title=filename or "DOCX Content",
                content=text,
                page_number=1,
                section_type="text"
            )
        ]
