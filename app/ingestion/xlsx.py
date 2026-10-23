import io
from typing import List, Optional
import openpyxl
from app.core.logging import logger
from app.ingestion.base import BaseParser
from app.schemas.document import ParsedSection


class XLSXParser(BaseParser):
    """Structured spreadsheet parser converting sheets and rows into searchable tabular contexts."""

    def parse_bytes(self, content: bytes, filename: str) -> List[ParsedSection]:
        sections: List[ParsedSection] = []
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)

            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                rows = list(sheet.iter_rows(values_only=True))

                if not rows:
                    continue

                # Filter out completely empty rows
                non_empty_rows = [r for r in rows if any(cell is not None and str(cell).strip() != "" for cell in r)]
                if not non_empty_rows:
                    continue

                # First non-empty row as header
                headers = [str(cell).strip() if cell is not None else f"Column_{i+1}" for i, cell in enumerate(non_empty_rows[0])]

                # Build both markdown table and key-value narrative representations
                table_lines = [
                    "| " + " | ".join(headers) + " |",
                    "| " + " | ".join(["---"] * len(headers)) + " |"
                ]

                row_narratives = []
                for row_idx, row in enumerate(non_empty_rows[1:], start=2):
                    row_vals = [str(cell).strip() if cell is not None else "" for cell in row]
                    table_lines.append("| " + " | ".join(row_vals) + " |")

                    # Key-value representation: e.g. "Row 2: Product: A | Price: 100 | Stock: 10"
                    kv_pairs = [f"{h}: {val}" for h, val in zip(headers, row_vals) if val]
                    if kv_pairs:
                        row_narratives.append(f"Row {row_idx}: " + " | ".join(kv_pairs))

                # Combine markdown table with explicit record sentences for high semantic retrieval recall
                content_str = (
                    f"### Sheet: {sheet_name}\n\n"
                    + "#### Summary Table:\n"
                    + "\n".join(table_lines)
                    + "\n\n#### Individual Records:\n"
                    + "\n".join(row_narratives)
                )

                sections.append(
                    ParsedSection(
                        title=f"Spreadsheet: {sheet_name}",
                        content=content_str,
                        page_number=1,
                        section_type="table",
                        metadata={
                            "filename": filename,
                            "sheet_name": sheet_name,
                            "row_count": len(non_empty_rows)
                        }
                    )
                )

        except Exception as e:
            logger.error(f"Failed to parse XLSX {filename}: {e}")
            raise e

        return sections

    def parse_text(self, text: str, filename: Optional[str] = None) -> List[ParsedSection]:
        return [
            ParsedSection(
                title=filename or "Spreadsheet Content",
                content=text,
                page_number=1,
                section_type="table"
            )
        ]
